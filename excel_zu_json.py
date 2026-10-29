#!/usr/bin/env python3
"""Excel-Arbeitsmappe (alle Tabellenblaetter) in eine JSON-Datei exportieren.

Annahme: Pro Tabellenblatt ein Datenblock mit einer Spaltenkopfzeile;
alle folgenden nichtleeren Zeilen sind Datensaetze.

Aufruf: py excel_zu_json.py "C:\\Pfad\\DZIF_Master.xlsx"
Ohne Dateipfad erscheint ein Dateiauswahldialog (falls tkinter verfuegbar).
"""

from __future__ import annotations

import argparse
from datetime import date, datetime, time, timedelta
from decimal import Decimal
import json
from pathlib import Path
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("FEHLER: openpyxl ist nicht installiert. Befehl: py -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)


def as_json(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (str, bool, int, float)):
        return value
    return str(value)


def find_file():
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        filename = filedialog.askopenfilename(
            title="Excel-Masterdatei auswaehlen",
            filetypes=[("Excel-Dateien", "*.xlsx *.xlsm"), ("Alle Dateien", "*.*")],
        )
        root.destroy()
        return Path(filename) if filename else None
    except Exception:
        return None


def nonempty(value):
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def unique_headers(row):
    seen = set()
    result = []
    for i, value in enumerate(row, start=1):
        base = str(value).strip() if nonempty(value) else f"Spalte_{i}"
        label = base
        count = 2
        while label in seen:
            label = f"{base}_{count}"
            count += 1
        seen.add(label)
        result.append(label)
    return result


def convert(src: Path, dst: Path):
    # data_only liest den letzten von Excel gespeicherten Formelwert, nicht den Formeltext.
    wb = load_workbook(src, read_only=True, data_only=True)
    all_sheets = {}
    for ws in wb.worksheets:
        # Bei manchen Exporten ist der 'used range' falsch hinterlegt.
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        headers = None
        output = []
        for raw_row in rows:
            if not any(nonempty(x) for x in raw_row):
                continue
            if headers is None:
                headers = unique_headers(raw_row)
                continue
            # Manche Datenzeilen gehen ueber die Breite des ersten Headers hinaus.
            if len(raw_row) > len(headers) and any(nonempty(x) for x in raw_row[len(headers):]):
                headers.extend(f"Spalte_{j}" for j in range(len(headers) + 1, len(raw_row) + 1))
            record = {
                key: as_json(raw_row[i]) if i < len(raw_row) else None
                for i, key in enumerate(headers)
            }
            output.append(record)
        all_sheets[ws.title] = output
        print(f"  {ws.title}: {len(output)} Datensaetze")
    wb.close()
    with dst.open("w", encoding="utf-8") as f:
        json.dump(all_sheets, f, ensure_ascii=False, indent=2, allow_nan=False)
        f.write("\n")
    # Kontrollieren, dass ein valides JSON-Dokument geschrieben wurde.
    with dst.open("r", encoding="utf-8") as f:
        parsed = json.load(f)
    print(f"\nFERTIG: {dst}")
    print(f"Tabellenblaetter: {len(parsed)} | JSON-Dateigroesse: {dst.stat().st_size / 1024:.1f} KB")


def main():
    parser = argparse.ArgumentParser(description="Alle Excel-Tabellenblaetter in eine JSON-Datei exportieren.")
    parser.add_argument("excel", nargs="?", help="Pfad zur .xlsx- oder .xlsm-Datei")
    args = parser.parse_args()
    source = Path(args.excel) if args.excel else find_file()
    if source is None:
        print("Keine Datei ausgewaehlt. Tipp: py excel_zu_json.py DATEINAME.xlsx", file=sys.stderr)
        return 2
    if not source.is_file() or source.suffix.lower() not in {".xlsx", ".xlsm"}:
        print(f"FEHLER: Kein gueltiger Excel-Dateipfad: {source}", file=sys.stderr)
        return 2
    target = source.with_suffix(".json")
    if target.exists():
        answer = input(f"{target.name} existiert bereits. Ueberschreiben? [j/N]: ").strip().lower()
        if answer not in {"j", "ja", "y", "yes"}:
            print("Abgebrochen. Vorhandene JSON-Datei wurde nicht veraendert.")
            return 0
    try:
        convert(source, target)
    except Exception as exc:
        print(f"FEHLER: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
